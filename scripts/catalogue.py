#!/usr/bin/env python3
"""
Art Studio price catalogue: a spreadsheet the owner edits, and the import that puts its prices on the site.

    npm run catalogue:export             python3 -I scripts/catalogue.py export [--force]
    npm run catalogue:import             python3 -I scripts/catalogue.py import [file.xlsx] [--dry-run] [--no-export]

export  writes catalogue/Art-Studio-Price-Catalogue.xlsx from src/data/prices.json (prices), src/config/prices.ts
        (codes, sizes) and src/data/seed.json (names, photos). It refuses to replace a catalogue that holds edits
        not imported yet, unless given --force.
import  reads the catalogue (or the file given), checks every row, prints each change (old → new) and writes
        src/data/prices.json. It writes nothing if anything is wrong, and only after src/config/prices.ts has
        loaded the new prices without complaint. It never touches TypeScript files. Afterwards it rebuilds the
        catalogue so it shows the saved prices, unless given --no-export or the catalogue is open in Excel.

Rows are matched by the grey Key column: "product/<slug>", "product/<slug>/<size id>", "product/<slug>/lights",
"category/<slug>" and "extra/<name>". Deleting rows takes things off the site:
  every row of a product   the product (prices.json keeps its entry with "removed": true, so its code is never
                           reused; catalog.ts hides it; delete the marker to bring it back)
  one size row             that size ("removed": true on the size)
  the lights row           the LED lights add-on for that product ("removed": true on the lights)
  a category row           its "From" price becomes "price on request" (amount null)
  a checkout charge row    nothing: these can't be removed, so the price stays
A tab whose rows are all gone and whose name is missing from the file is refused (probably the wrong file).

Needs openpyxl and Pillow, and Node with the project's node_modules (to read src/config/prices.ts).
"""

import argparse
import copy
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date
from pathlib import Path

try:
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.worksheet.datavalidation import DataValidation
    from PIL import Image as PILImage
except ImportError as error:
    sys.exit(f"scripts/catalogue.py needs openpyxl and Pillow ({error}).")

ROOT = Path(__file__).resolve().parent.parent
PRICES_JSON = ROOT / "src" / "data" / "prices.json"
SEED_JSON = ROOT / "src" / "data" / "seed.json"
PUBLIC = ROOT / "public"
CATALOGUE = ROOT / "catalogue" / "Art-Studio-Price-Catalogue.xlsx"

# Diwali and Car Gifts are priced 20% below our market estimate (owner, 29–30 Oct 2026).
MARKET_SHARE = (4, 5)


def below_market(market):
    """80% of the market estimate, rounded to the nearest rupee (as Math.round did in prices.ts)."""
    num, den = MARKET_SHARE
    return (2 * market * num + den) // (2 * den)


# ---------------------------------------------------------------------------------------------------------------
# Reading the price list: prices.json, prices.ts (through Node) and seed.json
# ---------------------------------------------------------------------------------------------------------------

# Loads src/config/prices.ts and src/config/launch.ts with the project's TypeScript compiler (transpile only) and
# prints what the site sees. argv: repo root, prices.json to use. prices.ts throws on a bad prices.json.
NODE_READER = r"""
const fs = require("fs"), path = require("path");
const [root, pricesFile] = process.argv.slice(1);
const ts = require(path.join(root, "node_modules", "typescript"));
const loaded = new Map();
function load(file) {
  if (file.endsWith(".json")) return JSON.parse(fs.readFileSync(file, "utf8"));
  if (loaded.has(file)) return loaded.get(file).exports;
  const js = ts.transpileModule(fs.readFileSync(file, "utf8"), {
    fileName: file,
    compilerOptions: { module: ts.ModuleKind.CommonJS, target: ts.ScriptTarget.ES2019, esModuleInterop: true },
  }).outputText;
  const mod = { exports: {} };
  loaded.set(file, mod);
  const req = (spec) => {
    if (spec === "@/data/prices.json") return load(pricesFile);
    const base = spec.startsWith("@/") ? path.join(root, "src", spec.slice(2)) : spec.startsWith(".") ? path.resolve(path.dirname(file), spec) : null;
    if (!base) throw new Error(`${path.relative(root, file)} imports "${spec}", which scripts/catalogue.py can't load`);
    for (const f of [base, base + ".ts", base + ".json", path.join(base, "index.ts")]) if (fs.existsSync(f) && fs.statSync(f).isFile()) return load(f);
    throw new Error(`can't find "${spec}" (imported by ${path.relative(root, file)})`);
  };
  new Function("exports", "require", "module", js)(mod.exports, req, mod);
  return mod.exports;
}
try {
  const prices = load(path.join(root, "src", "config", "prices.ts"));
  const { launch } = load(path.join(root, "src", "config", "launch.ts"));
  process.stdout.write(JSON.stringify({
    productPrices: prices.productPrices,
    removedProducts: Array.from(prices.removedProducts || []),
    customCategoryPrices: prices.customCategoryPrices,
    extraPrices: prices.extraPrices,
    productLines: prices.productLines,
    lightsSuffix: prices.LIGHTS_CODE_SUFFIX,
    lightsOnSale: Boolean(launch && launch.lights),
  }));
} catch (error) {
  process.stderr.write(String((error && error.message) || error));
  process.exit(1);
}
"""


class Stop(Exception):
    """Something the person running the script has to fix; printed without a traceback."""


def read_price_list(prices_file=None):
    """What the site sees: src/config/prices.ts loaded with the given prices.json (default: the project's)."""
    prices_file = prices_file or PRICES_JSON
    node = shutil.which("node")
    if not node:
        raise Stop("Node.js is needed to read src/config/prices.ts (install it, or run this from the project's terminal).")
    if not (ROOT / "node_modules" / "typescript").is_dir():
        raise Stop("node_modules/typescript is missing: run npm install first.")
    result = subprocess.run([node, "-e", NODE_READER, str(ROOT), str(prices_file)], cwd=ROOT, capture_output=True, text=True, timeout=120)
    if result.returncode != 0:
        raise Stop(f"src/config/prices.ts did not accept the prices:\n{result.stderr.strip()}")
    return json.loads(result.stdout)


def read_json(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def is_removed(value):
    """A product, size or lights entry taken off the site (its rows were deleted from the catalogue)."""
    return isinstance(value, dict) and value.get("removed") is True


def product_slots(slug, product):
    """{key: price dict} for one product's rows, leaving out what was taken off the site."""
    if is_removed(product):
        return {}
    slots = {}
    for field in ("price", "startingPrice"):
        if field in product:
            slots[f"product/{slug}"] = product[field]
    for size_id, price in (product.get("sizes") or {}).items():
        if not is_removed(price):
            slots[f"product/{slug}/{size_id}"] = price
    if product.get("lights") and not is_removed(product["lights"]):
        slots[f"product/{slug}/lights"] = product["lights"]
    return slots


def price_slots(data):
    """Every editable price in prices.json, by catalogue key. The values are the dicts inside `data`."""
    slots = {}
    for slug, product in data["products"].items():
        slots.update(product_slots(slug, product))
    for slug, price in data["customCategories"].items():
        slots[f"category/{slug}"] = price
    for name, price in data["extras"].items():
        slots[f"extra/{name}"] = price
    return slots


def site_slots(price_list):
    """The same keys, from what prices.ts produced: {key: (amount, confirmed)}."""
    pair = lambda p: (p["amount"], p["confirmed"])  # noqa: E731
    slots = {}
    for slug, p in price_list["productPrices"].items():
        if p["kind"] == "sizes":
            for size in p["sizes"]:
                slots[f"product/{slug}/{size['id']}"] = pair(size["price"])
            if p.get("lights"):
                slots[f"product/{slug}/lights"] = pair(p["lights"])
        else:
            slots[f"product/{slug}"] = pair(p["price"] if p["kind"] == "fixed" else p["startingPrice"])
    for slug, p in price_list["customCategoryPrices"].items():
        slots[f"category/{slug}"] = pair(p)
    for name, p in price_list["extraPrices"].items():
        slots[f"extra/{name}"] = pair(p)
    return slots


def _inline(value):
    if isinstance(value, dict):
        return "{ " + ", ".join(f"{json.dumps(k, ensure_ascii=False)}: {_inline(v)}" for k, v in value.items()) + " }" if value else "{}"
    if isinstance(value, list):
        return "[" + ", ".join(_inline(v) for v in value) + "]"
    return json.dumps(value, ensure_ascii=False)


def _depth(value):
    if isinstance(value, dict):
        return 1 + max((_depth(v) for v in value.values()), default=0)
    if isinstance(value, list):
        return 1 + max((_depth(v) for v in value), default=0)
    return 0


def to_json(value, level=0, used=0):
    """JSON with one price per line, so changes read well in a diff."""
    one_line = _inline(value)
    if not isinstance(value, (dict, list)) or (_depth(value) <= 2 and used + len(one_line) <= 110):
        return one_line
    pad, inner = "  " * level, "  " * (level + 1)
    if isinstance(value, list):
        return "[\n" + ",\n".join(inner + to_json(v, level + 1, len(inner)) for v in value) + "\n" + pad + "]"
    items = []
    for k, v in value.items():
        head = f"{inner}{json.dumps(k, ensure_ascii=False)}: "
        items.append(head + to_json(v, level + 1, len(head)))
    return "{\n" + ",\n".join(items) + "\n" + pad + "}"


def inr(amount, empty="price on request"):
    """₹1,23,456 (Indian grouping)."""
    if amount is None:
        return empty
    digits = str(abs(amount))
    head, tail = digits[:-3], digits[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ("-" if amount < 0 else "") + "₹" + ",".join(groups + [tail])


# ---------------------------------------------------------------------------------------------------------------
# The catalogue rows
# ---------------------------------------------------------------------------------------------------------------

PRODUCT_SHEETS = ["Retirement Gifts", "Diwali", "Christmas", "Car Gifts", "Shop", "Made to order"]
SHEET_BY_CATEGORY = {
    "retirement-tribute-frames": "Retirement Gifts",
    "diwali-decor-gifts": "Diwali",
    "christmas-decor-gifts": "Christmas",
    "car-gifts": "Car Gifts",
}
EXTRAS_SHEET = "Categories & extras"
HOWTO_SHEET = "How to use"
TAB_COLOURS = {"Retirement Gifts": "5A3E2B", "Diwali": "D5851F", "Christmas": "87A96B", "Car Gifts": "5A3E2B", "Shop": "D5851F", "Made to order": "87A96B", EXTRAS_SHEET: "5A3E2B"}

EXTRAS = {
    "freeShippingFrom": ("Free shipping from", "Orders at or above this total ship free."),
    "shippingFee": ("Delivery fee", "Charged on orders below the free-shipping total. 0 = delivery is always free."),
    "giftWrap": ("Gift wrap", "Per order, when the customer asks for it. 0 = free."),
    "codFee": ("Cash on Delivery fee", "Added when the customer pays on delivery. 0 = no fee."),
}


def short_name(name):
    return name.split(" | ")[0]


def photo_file(product):
    """The product's main photo: the 4:5 tile (or frame photo) for tributes, else the first image."""
    if not product:
        return None
    gallery = product.get("gallery") or {}
    for src in (gallery.get("tile"), gallery.get("frame"), *(product.get("images") or [])[:1]):
        if src and (PUBLIC / src.lstrip("/")).is_file():
            return PUBLIC / src.lstrip("/")
    return None


def variant_text(variant_codes):
    """('KAAL + letter', 'Any letter A–Z') style descriptions of a product's choices."""
    names = list(variant_codes)
    if names == [chr(c) for c in range(ord("A"), ord("Z") + 1)] and all(variant_codes[n] == n for n in names):
        return "+ letter (A–Z)", "Any letter A–Z, same price"
    return "+ " + " / ".join(variant_codes.values()), "Same price for: " + ", ".join(names)


def build_rows(data, price_list, seed):
    """Every catalogue row, in sheet order. Values come from prices.json; codes and sizes from prices.ts."""
    products = {p["slug"]: p for p in seed["products"]}
    categories = {c["slug"]: c for c in seed["categories"]}
    lines = price_list["productLines"]
    slots = price_slots(data)
    rows = []

    def row(key, **fields):
        slot = slots.get(key, {"amount": None, "confirmed": False})
        rows.append({"key": key, "amount": slot["amount"], "confirmed": slot["confirmed"], "market": slot.get("market"), "first": False, **fields})

    for slug, entry in price_list["productPrices"].items():
        product = products.get(slug)
        saved = data["products"].get(slug, {})
        category = product["category"] if product else (lines.get(entry["code"][:2]) or {}).get("category")
        sheet = SHEET_BY_CATEGORY.get(category) or ("Made to order" if entry["kind"] == "from" else "Shop")
        name = product["name"] if product else slug
        base = {"sheet": sheet, "slug": slug, "name": name, "photo": photo_file(product), "product_label": f"{short_name(name)} ({entry['code']})"}
        notes = [saved.get("note")]
        if not product:
            notes.append("Not in the product list (seed.json), so not on the site.")
        elif not product.get("published", True):
            notes.append("Hidden on the site for now.")
        start = len(rows)
        if entry["kind"] == "sizes":
            for size in entry["sizes"]:
                row(f"product/{slug}/{size['id']}", code=size["code"], option=size["label"], label=f"{short_name(name)} · {size['label']} ({size['code']})", **base)
            if entry.get("lights"):
                example = entry["sizes"][0]["code"] + price_list["lightsSuffix"] if entry["sizes"] else ""
                lights_note = "Added to the size price."
                if not price_list["lightsOnSale"]:
                    lights_note += " Lights are switched off on the site for now; this price is kept for later."
                row(f"product/{slug}/lights", code=f"+ {price_list['lightsSuffix']} (e.g. {example})", option="LED lights add-on", note=lights_note, label=f"{short_name(name)} · LED lights", **base)
        else:
            code, option = entry["code"], ("One price" if entry["kind"] == "fixed" else "Starting price (“From ₹…”)")
            if entry.get("variantCodes"):
                more, option = variant_text(entry["variantCodes"])
                code = f"{entry['code']} {more}"
                if entry["kind"] == "from":
                    option = option.replace("Same price for", "Starting price for")
            row(f"product/{slug}", code=code, option=option, label=f"{short_name(name)} ({entry['code']})", **base)
        rows[start]["first"] = True
        for r in rows[start:]:
            row_notes = [n for n in notes if n] if r is rows[start] else []
            if r.get("note"):
                row_notes.append(r["note"])
            if r["market"] is not None and r["amount"] is not None:
                if r["amount"] == below_market(r["market"]):
                    row_notes.append("20% below the market estimate.")
                else:
                    row_notes.append(f"Price set by hand (20% below the estimate would be {inr(below_market(r['market']))}).")
            if r["amount"] is None:
                row_notes.append("Empty: no price shown (price on request)." if entry["kind"] == "from" else "Empty: price on request; can't be added to the cart.")
            r["note"] = " ".join(row_notes)

    for slug in data["customCategories"]:
        category = categories.get(slug)
        notes = []
        if not category:
            notes.append("Not in the category list (seed.json).")
        elif category.get("status") == "coming-soon":
            notes.append("Coming soon on the site.")
        if slots[f"category/{slug}"]["amount"] is None:
            notes.append("Empty: “Price on request”.")
        name = category["name"] if category else slug
        row(f"category/{slug}", sheet=EXTRAS_SHEET, section="Made to order: “From” price", name=name, label=f"{name} (“From” price)", note=" ".join(notes))
    for name in data["extras"]:
        label, note = EXTRAS.get(name, (name, ""))
        row(f"extra/{name}", sheet=EXTRAS_SHEET, section="Checkout charges", name=label, label=f"{label} (checkout)", note=note)
    return rows


def labels(rows):
    """Readable name for each key, for the import summary, plus "product/<slug>" for every product."""
    names = {r["key"]: r["label"] for r in rows}
    for r in rows:
        if r.get("product_label"):
            names.setdefault(f"product/{r['slug']}", r["product_label"])
    return names


# ---------------------------------------------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------------------------------------------

BROWN, AMBER, SAGE, INK = "5A3E2B", "D5851F", "87A96B", "1E140D"
HEADER_FILL = PatternFill("solid", fgColor=BROWN)
EDIT_HEADER_FILL = PatternFill("solid", fgColor=AMBER)
EDIT_FILL = PatternFill("solid", fgColor="FCEBD0")
KEY_HEADER_FILL = PatternFill("solid", fgColor="BDBDBD")
KEY_FILL = PatternFill("solid", fgColor="EDEDED")
SECTION_FILL = PatternFill("solid", fgColor="E3EBDC")
HEADER_FONT = Font(bold=True, color="FFFFFF")
EDIT_HEADER_FONT = Font(bold=True, color=INK)
KEY_FONT = Font(color="8C8C8C", size=9)
NOTE_FONT = Font(color="5C4A3D", size=10)
MUTED_FONT = Font(color="8C8C8C", italic=True, size=9)
PRODUCT_TOP = Border(top=Side(style="thin", color="D8CBBE"))
WRAP = Alignment(wrap_text=True, vertical="center")
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
PHOTO_BOX = (92, 100)  # thumbnail fits in this many pixels (width, height); stored at 2× for sharp screens
PHOTO_ROW_PT = 84  # a product's rows add up to at least this height in points, so its photo fits
MONEY = "#,##0"


def thumbnail(path):
    try:
        with PILImage.open(path) as image:
            image = image.convert("RGB")
            scale = min(PHOTO_BOX[0] / image.width, PHOTO_BOX[1] / image.height)
            width, height = max(1, round(image.width * scale)), max(1, round(image.height * scale))
            image = image.resize((width * 2, height * 2), PILImage.LANCZOS)
            data = io.BytesIO()
            image.save(data, "JPEG", quality=82, optimize=True)
            data.seek(0)
            return data, width, height
    except OSError:
        return None


def add_photo(ws, path, column, row):
    thumb = thumbnail(path)
    if not thumb:
        return False
    data, width, height = thumb
    image = XLImage(data)
    image.width, image.height = width, height
    marker = AnchorMarker(col=column - 1, colOff=pixels_to_EMU(6), row=row - 1, rowOff=pixels_to_EMU(4))
    image.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(width), pixels_to_EMU(height)))
    ws.add_image(image)
    return True


def text_lines(text, width):
    """Rough number of wrapped lines `text` needs in a column `width` characters wide."""
    if not text:
        return 1
    usable = max(4, int(width * 1.15) - 2)
    return sum(max(1, -(-len(part) // usable)) for part in str(text).split("\n"))


def write_header(ws, columns):
    for i, (title, width, kind) in enumerate(columns, start=1):
        cell = ws.cell(1, i, title)
        if kind in ("amount", "confirmed", "market"):
            cell.fill, cell.font = EDIT_HEADER_FILL, EDIT_HEADER_FONT
        elif kind == "key":
            cell.fill, cell.font = KEY_HEADER_FILL, Font(bold=True, color="555555")
        else:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = CENTRE
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def add_validations(ws, columns, last_row, min_price):
    if last_row < 2:
        return
    col = {kind: get_column_letter(i) for i, (_, _, kind) in enumerate(columns, start=1)}
    confirmed = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True, errorTitle="Confirmed", error="Choose Yes or No.")
    confirmed.add(f"{col['confirmed']}2:{col['confirmed']}{last_row}")
    ws.add_data_validation(confirmed)
    price_rule = "a whole number above 0, or leave it empty" if min_price else "a whole number, 0 or more"
    price = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=str(min_price), allow_blank=True, showErrorMessage=True, errorTitle="Price", error=f"Type the price in whole rupees: {price_rule}.")
    price.add(f"{col['amount']}2:{col['amount']}{last_row}")
    ws.add_data_validation(price)
    if "market" in col:
        market = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True, showErrorMessage=True, errorTitle="Market estimate", error="Type the estimate in whole rupees.")
        market.add(f"{col['market']}2:{col['market']}{last_row}")
        ws.add_data_validation(market)


def fill_cells(ws, r, columns, row, first_of_block):
    for i, (_, _, kind) in enumerate(columns, start=1):
        cell = ws.cell(r, i)
        if kind == "name":
            cell.value = row["name"] if first_of_block else None
            cell.font, cell.alignment = Font(bold=True, color=INK), WRAP
        elif kind in ("code", "option", "section"):
            cell.value = row.get(kind)
            cell.alignment = WRAP
        elif kind == "amount":
            cell.value, cell.number_format, cell.fill, cell.alignment = row["amount"], MONEY, EDIT_FILL, CENTRE
        elif kind == "confirmed":
            cell.value, cell.fill, cell.alignment = "Yes" if row["confirmed"] else "No", EDIT_FILL, CENTRE
        elif kind == "market":
            cell.value, cell.number_format, cell.fill, cell.alignment = row["market"], MONEY, EDIT_FILL, CENTRE
        elif kind == "note":
            cell.value, cell.font, cell.alignment = row["note"] or None, NOTE_FONT, WRAP
        elif kind == "key":
            cell.value, cell.fill, cell.font, cell.alignment = row["key"], KEY_FILL, KEY_FONT, WRAP
        if first_of_block and r > 2:
            cell.border = PRODUCT_TOP


def row_height(columns, shown):
    """Height in points for a row showing these values (wrapped text gets more lines)."""
    lines = max(text_lines(shown.get(kind), width) for _, width, kind in columns if kind in ("name", "code", "option", "note", "section"))
    return max(20, lines * 13.5 + 6)


def write_product_sheet(ws, rows):
    has_market = any(r["market"] is not None for r in rows)
    columns = [("Photo", 14, "photo"), ("Product", 30, "name"), ("Code", 16, "code"), ("Option", 24, "option"), ("Price ₹", 11, "amount"), ("Confirmed", 11, "confirmed")]
    if has_market:
        columns.append(("Market estimate ₹", 13, "market"))
    columns += [("Notes", 46, "note"), ("Key (don't edit)", 36, "key")]
    write_header(ws, columns)
    photos = 0
    r = 2
    blocks = []
    for row in rows:
        if row["first"]:
            blocks.append([])
        blocks[-1].append(row)
    for block in blocks:
        heights = [row_height(columns, {**row, "name": row["name"] if i == 0 else None}) for i, row in enumerate(block)]
        spare = PHOTO_ROW_PT - sum(heights)
        if spare > 0:
            heights = [h + spare / len(heights) for h in heights]
        for i, row in enumerate(block):
            fill_cells(ws, r + i, columns, row, i == 0)
            ws.row_dimensions[r + i].height = round(heights[i], 1)
        if block[0]["photo"] and add_photo(ws, block[0]["photo"], 1, r):
            photos += 1
        else:
            cell = ws.cell(r, 1, "No photo yet")
            cell.font, cell.alignment = MUTED_FONT, CENTRE
        r += len(block)
    add_validations(ws, columns, r - 1, 1)
    return photos


def write_extras_sheet(ws, rows):
    columns = [("Section", 22, "section"), ("Item", 32, "name"), ("Price ₹", 11, "amount"), ("Confirmed", 11, "confirmed"), ("Notes", 52, "note"), ("Key (don't edit)", 32, "key")]
    write_header(ws, columns)
    previous = None
    for r, row in enumerate(rows, start=2):
        shown = {**row, "section": row["section"] if row["section"] != previous else None}
        fill_cells(ws, r, columns, shown, row["section"] != previous)
        ws.cell(r, 2).value, ws.cell(r, 2).font = row["name"], Font(color=INK)
        ws.row_dimensions[r].height = row_height(columns, shown)
        previous = row["section"]
    # Checkout charges may be 0; the import checks which rows need more.
    add_validations(ws, columns, len(rows) + 1, 0)


def write_howto(ws, rows, lights_on_sale):
    today = date.today()
    counts = {sheet: len({r["slug"] for r in rows if r["sheet"] == sheet}) for sheet in PRODUCT_SHEETS}
    tabs = [f"{sheet}: {n} {'piece' if n == 1 else 'pieces'}" for sheet, n in counts.items() if n]
    tabs.append(f"{EXTRAS_SHEET}: the “From” price on each made-to-order category, and the checkout charges (delivery, gift wrap, COD fee)")
    lines = [
        ("Art Studio · Price catalogue", "title"),
        (f"Prices as on {today.day} {today:%b %Y}. All prices are in rupees, taxes included.", "muted"),
        ("", None),
        ("How to change prices", "heading"),
        ("1. Change only the orange columns: Price ₹, Confirmed, and (Diwali, Car Gifts) Market estimate ₹. The other columns are for reference; changes there are ignored.", None),
        ("2. Type prices in whole rupees, like 2499. Leave a price empty for “price on request” (that piece then can't be added to the cart). Checkout charges can be 0 (no charge) but not empty.", None),
        ("3. Confirmed: choose Yes when a price is final, No while it is still a guess.", None),
        ("4. Retirement Gifts have one row per size and one for the LED lights add-on, which is added to the size price."
         + ("" if lights_on_sale else " Lights are switched off on the site for now; their price is kept for later."), None),
        ("5. Diwali and Car Gifts are priced 20% below our market estimate. Change just the Market estimate and the price is worked out for you when you import (80% of it). Change the Price and that price is used as it is.", None),
        ("6. Don't change the grey Key column: it tells the site which price is which. New products are added by the developer first; then export a fresh copy.", None),
        ("7. To take something off the site, delete its rows. All of a piece's rows: the piece comes off the site. One size row: that size is no longer offered. The LED lights row: no lights for that piece. A category's “From” row: it shows “Price on request”. Checkout charges can't be deleted (type 0 for no charge). The import lists everything it takes off; to bring something back, ask the developer.", None),
        ("", None),
        ("Putting the prices on the site", "heading"),
        ("1. Save this file (keep the .xlsx format and the name) and close it.", None),
        ("2. In the project folder, run:  npm run catalogue:import", "code"),
        ("   It lists every change (old → new). If something looks wrong it says what and where, and changes nothing. After importing it rebuilds this file with the saved prices.", None),
        ("3. Check the prices on the site (npm run preview), then commit and push to update the live site.", None),
        ("", None),
        ("A fresh copy of this file:  npm run catalogue:export", "code"),
        ("It won't replace a file that has edits not yet imported.", None),
        ("", None),
        ("Tabs", "heading"),
        *[(f"• {t}", None) for t in tabs],
        ("Retirement Gifts' “From” price is their lowest size price, so it has no row of its own.", "muted"),
    ]
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    styles = {
        "title": Font(bold=True, size=18, color=BROWN),
        "heading": Font(bold=True, size=13, color=BROWN),
        "muted": Font(italic=True, color="7A6A5D"),
        "code": Font(bold=True, color=INK),
    }
    for r, (text, style) in enumerate(lines, start=1):
        cell = ws.cell(r, 2, text or None)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if style:
            cell.font = styles[style]
        ws.row_dimensions[r].height = 30 if style == "title" else max(18, text_lines(text, 110) * 15 + 3)
    ws.sheet_view.showGridLines = False


def excel_lock(path):
    """The lock file Excel keeps next to a workbook it has open (~$name.xlsx), if there is one."""
    for name in (f"~${path.name}", f"~${path.name[2:]}"):
        if path.with_name(name).exists():
            return path.with_name(name)
    return None


def write_catalogue(path, data, price_list, seed):
    lock = excel_lock(path)
    if lock:
        raise Stop(f"{display(path)} is open in Excel, so it was left as it is. Close it and try again (if Excel is closed already, delete {display(lock)} first).")
    rows = build_rows(data, price_list, seed)
    wb = openpyxl.Workbook()
    write_howto(wb.active, rows, price_list["lightsOnSale"])
    wb.active.title = HOWTO_SHEET
    photos = 0
    for sheet in PRODUCT_SHEETS:
        sheet_rows = [r for r in rows if r["sheet"] == sheet]
        if sheet_rows:
            ws = wb.create_sheet(sheet)
            ws.sheet_properties.tabColor = TAB_COLOURS[sheet]
            photos += write_product_sheet(ws, sheet_rows)
    ws = wb.create_sheet(EXTRAS_SHEET)
    ws.sheet_properties.tabColor = TAB_COLOURS[EXTRAS_SHEET]
    write_extras_sheet(ws, [r for r in rows if r["sheet"] == EXTRAS_SHEET])
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.stem}.saving.xlsx")
    try:
        wb.save(temp)
        os.replace(temp, path)
    except PermissionError:
        if temp.exists():
            temp.unlink()
        raise Stop(f"Couldn't write {display(path)}: is it open in Excel? Close it and try again.")
    pieces = len({r["slug"] for r in rows if r["sheet"] != EXTRAS_SHEET})
    return f"{len(rows)} prices for {pieces} pieces, {photos} photos"


def unpriced_products(seed, price_list):
    known = set(price_list["productPrices"]) | set(price_list["removedProducts"])
    return [p["slug"] for p in seed["products"] if p["slug"] not in known]


def display(path):
    try:
        return str(Path(path).resolve().relative_to(ROOT))
    except ValueError:
        return str(path)


def command_export(force):
    data, seed = read_json(PRICES_JSON), read_json(SEED_JSON)
    price_list = read_price_list()
    if CATALOGUE.exists() and not force:
        errors, changes, removals = compare_workbook(CATALOGUE, data)
        edits = len(changes) + len(removals)
        if errors or edits:
            found = f"{edits} edit{'s' if edits != 1 else ''} not imported yet" if edits else "rows that can't be read"
            raise Stop(
                f"{display(CATALOGUE)} has {found}, so it was left as it is.\n"
                "Import it first (npm run catalogue:import), or run npm run catalogue:export -- --force to replace it and lose those edits."
            )
    summary = write_catalogue(CATALOGUE, data, price_list, seed)
    print(f"Wrote {display(CATALOGUE)}: {summary}.")
    missing = unpriced_products(seed, price_list)
    if missing:
        print(f"Not in the catalogue (no entry in src/config/prices.ts): {', '.join(missing)}")
    if price_list["removedProducts"]:
        print(f"Taken off the site in an earlier import, so not in the catalogue: {', '.join(price_list['removedProducts'])}")


# ---------------------------------------------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------------------------------------------

ABSENT = object()  # a column this sheet doesn't have


class Bad(Exception):
    pass


def header_columns(ws):
    """(header row, {kind: column}) for a catalogue sheet, or (None, {}) if the sheet has no Key column."""
    for r in range(1, min(ws.max_row, 5) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            title = str(ws.cell(r, c).value or "").strip().lower()
            for kind in ("key", "price", "confirmed", "market"):
                if title.startswith(kind) and kind not in found:
                    found[kind] = c
        if "key" in found:
            return r, found
    return None, {}


def read_workbook(path):
    """[{where, key, price, confirmed, market}] for every row with a key, problems found while reading, and the tab names."""
    try:
        values = openpyxl.load_workbook(path, data_only=True)
        formulas = openpyxl.load_workbook(path, data_only=False)
    except FileNotFoundError:
        raise Stop(f"{display(path)} not found. Run npm run catalogue:export to create it.")
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise Stop(f"Couldn't open {display(path)} as an Excel .xlsx file ({error}). Save it as .xlsx and try again.")
    entries, errors = [], []
    for ws in values.worksheets:
        header, cols = header_columns(ws)
        if not header:
            continue
        if "price" not in cols:
            errors.append(f"{ws.title}: no “Price ₹” column, so this tab can't be read.")
            continue
        fws = formulas[ws.title]

        def read(r, kind):
            if kind not in cols:
                return ABSENT
            value = ws.cell(r, cols[kind]).value
            formula = fws.cell(r, cols[kind]).value
            if value is None and isinstance(formula, str) and formula.startswith("="):
                raise Bad(f"{ws.title}, row {r}: the formula {formula} has no saved result. Type the number instead, or open and save the file in Excel.")
            return value

        for r in range(header + 1, ws.max_row + 1):
            where = f"{ws.title}, row {r}"
            key = str(ws.cell(r, cols["key"]).value or "").strip()
            try:
                price = read(r, "price")
                if not key:
                    if price not in (None, "", ABSENT):
                        errors.append(f"{where}: has a price but no key. Put the key back (export a fresh copy to see it), or delete the row.")
                    continue
                entries.append({"where": where, "key": key, "price": price, "confirmed": read(r, "confirmed"), "market": read(r, "market")})
            except Bad as error:
                errors.append(str(error))
    return entries, errors, values.sheetnames


def parse_rupees(value, where, what):
    """Whole rupees from a cell: 2499, 2499.0, "2,499", "₹2,499" or "Rs 2499". None when empty."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise Bad(f"{where}: {what} is {value}, not a number of rupees.")
    if isinstance(value, (int, float)):
        if float(value).is_integer():
            return int(value)
        raise Bad(f"{where}: {what} {value} is not whole rupees. Round it, e.g. {round(value)}.")
    if isinstance(value, str):
        text = re.sub(r"(?i)^\s*(₹|rs\.?|inr)\s*", "", value.strip()).replace(",", "").replace(" ", "")
        if re.fullmatch(r"-?\d+(\.0+)?", text):
            return int(float(text))
        raise Bad(f"{where}: {what} “{value}” is not a number of rupees.")
    raise Bad(f"{where}: {what} “{value}” is not a number of rupees.")


def parse_confirmed(value, where):
    if isinstance(value, bool):
        return value
    if value is None or isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    text = str(value).strip().lower()
    if text in ("yes", "y", "true", "1", "✓", "✔"):
        return True
    if text in ("no", "n", "false", "0", ""):
        return False
    raise Bad(f"{where}: Confirmed “{value}” should be Yes or No.")


def taken_off(data, key):
    """For a key that is no longer in the catalogue: what it was, if an earlier import took it off the site."""
    parts = key.split("/")
    product = data["products"].get(parts[1]) if parts[0] == "product" and len(parts) > 1 else None
    if not isinstance(product, dict):
        return None
    if is_removed(product):
        return "This piece"
    if len(parts) == 3 and parts[2] == "lights" and is_removed(product.get("lights")):
        return "The LED lights add-on"
    if len(parts) == 3 and is_removed((product.get("sizes") or {}).get(parts[2])):
        return "This size"
    return None


def remove_missing(new, old_slots, present, changes, errors, names):
    """
    Takes off the site, in `new`, whatever lost its rows (see the top of this file). Returns removals (keys of
    products, sizes and lights) and kept (checkout charges, which can't be removed). A deleted category row
    is added to `changes` as a price change to "price on request".
    """
    removals, kept = [], []
    for slug, product in new["products"].items():
        keys = list(product_slots(slug, product))
        missing = [k for k in keys if k not in present]
        if not missing:
            continue
        if len(missing) == len(keys):
            new["products"][slug] = {"removed": True, **product}
            removals.append(f"product/{slug}")
            continue
        sizes = product.get("sizes") or {}
        size_keys = [k for k in keys if not k.endswith("/lights")]
        if sizes and all(k in missing for k in size_keys):
            label = names.get(f"product/{slug}", slug)
            errors.append(f"{label}: every size row was deleted but the LED lights row is still there. Delete that row too to take the piece off the site, or put a size row back.")
            continue
        for key in missing:
            if key.endswith("/lights"):
                product["lights"] = {**product["lights"], "removed": True}
            else:
                size_id = key.rsplit("/", 1)[1]
                sizes[size_id] = {**sizes[size_id], "removed": True}
            removals.append(key)
    for slug, price in new["customCategories"].items():
        key = f"category/{slug}"
        if key not in present and price["amount"] is not None:
            # The owner chose "price on request" by deleting the row, so it counts as confirmed.
            price["amount"], price["confirmed"] = None, True
            changes.append((key, old_slots[key], dict(price), "deleted"))
    for name in new["extras"]:
        if f"extra/{name}" not in present:
            kept.append(f"extra/{name}")
    return removals, kept


def apply_entries(data, entries, names=None):
    """
    The new prices.json and what changed: (new, problems, changes, removals, kept). changes are price edits
    [(key, old, new, why)], `why` being "market" (worked out from a new market estimate), "deleted" (a category
    row was deleted) or None. removals and kept are keys of rows deleted from the sheet (see remove_missing).
    """
    names = names or {}
    new = copy.deepcopy(data)
    old_slots, slots = price_slots(data), price_slots(new)
    seen, errors, changes = {}, [], []
    for entry in entries:
        key, where = entry["key"], entry["where"]
        if key in seen:
            errors.append(f"{where}: the same key as {seen[key]} ({key}). Delete one of the two rows.")
            continue
        seen[key] = where
        if key not in slots:
            gone = taken_off(data, key)
            if gone:
                errors.append(f"{where}: {gone} was taken off the site in an earlier import. Delete this row (to bring it back, ask the developer).")
            else:
                errors.append(f"{where}: unknown key “{key}”. Keys must stay as exported; export a fresh copy if products have changed.")
            continue
        old, slot = old_slots[key], slots[key]
        is_extra = key.startswith("extra/")
        try:
            price = parse_rupees(entry["price"], where, "Price")
            confirmed = old["confirmed"] if entry["confirmed"] is ABSENT else parse_confirmed(entry["confirmed"], where)
            market = ABSENT if entry["market"] is ABSENT else parse_rupees(entry["market"], where, "Market estimate")
            if price is None and is_extra:
                raise Bad(f"{where}: Price can't be empty here. Use 0 for no charge.")
            if price is not None and price < (0 if is_extra else 1):
                raise Bad(f"{where}: Price {price} must be {'0 or more' if is_extra else 'above 0 (or empty for price on request)'}.")
            if market not in (None, ABSENT) and market < 1:
                raise Bad(f"{where}: Market estimate {market} must be above 0 (or empty).")
        except Bad as error:
            errors.append(str(error))
            continue
        amount, why = price, None
        if market is not ABSENT:
            if price == old["amount"] and market is not None and market != old.get("market"):
                amount, why = below_market(market), "market"  # only the estimate changed
            if market is None:
                slot.pop("market", None)
            else:
                slot["market"] = market
        slot["amount"], slot["confirmed"] = amount, confirmed
        if slot != old:
            changes.append((key, old, dict(slot), why))
    removals, kept = remove_missing(new, old_slots, set(seen), changes, errors, names)
    return new, errors, changes, removals, kept


def missing_tabs(rows, entries, titles):
    """Tabs the catalogue should have that are gone, rows and all: probably the wrong file, so refuse."""
    present = {e["key"] for e in entries}
    problems = []
    for sheet in dict.fromkeys(r["sheet"] for r in rows):
        if sheet not in titles and not any(r["key"] in present for r in rows if r["sheet"] == sheet):
            problems.append(
                f"The “{sheet}” tab is missing, so all of its prices would be taken off the site. Use the full catalogue"
                " (npm run catalogue:export makes a fresh copy). To take every piece on that tab off the site, keep the tab and delete its rows."
            )
    return problems


def compare_workbook(path, data):
    """Problems, price changes and removals in the workbook compared with prices.json."""
    entries, read_errors, _ = read_workbook(path)
    _, errors, changes, removals, _ = apply_entries(data, entries)
    return read_errors + errors, changes, removals


def describe(label, old, new, why):
    parts = []
    if old.get("market") != new.get("market"):
        parts.append(f"market estimate {inr(old.get('market'), 'none')} → {inr(new.get('market'), 'none')}")
    if old["amount"] != new["amount"]:
        reason = {"market": " (20% below the new estimate)", "deleted": " (its row was deleted)"}.get(why, "")
        parts.append(f"price {inr(old['amount'], 'empty')} → {inr(new['amount'], 'empty (price on request)')}{reason}")
    if old["confirmed"] != new["confirmed"]:
        parts.append(f"confirmed {'Yes' if old['confirmed'] else 'No'} → {'Yes' if new['confirmed'] else 'No'}")
    return f"  {label}: " + "; ".join(parts)


def print_list(title, items):
    if items:
        print(f"\n{title}")
        for item in items:
            print(f"  {item}")


def command_import(source, dry_run, no_export=False):
    source = Path(source).expanduser()
    if not source.is_absolute():
        source = Path(os.environ.get("INIT_CWD") or os.getcwd()) / source
    data, seed = read_json(PRICES_JSON), read_json(SEED_JSON)
    price_list = read_price_list()
    rows = build_rows(data, price_list, seed)
    names = labels(rows)
    entries, read_errors, titles = read_workbook(source)
    new, errors, changes, removals, kept = apply_entries(data, entries, names)
    errors = read_errors + missing_tabs(rows, entries, titles) + errors
    print(f"Read {display(source)}: {len(entries)} price rows.")
    if errors:
        print(f"\nNothing was changed. Fix {'this' if len(errors) == 1 else 'these'} in the spreadsheet and import again:")
        for error in errors:
            print(f"  • {error}")
        return 1
    kept_note = "Kept as they are (checkout charges can't be deleted; type 0 for no charge):"
    if not changes and not removals:
        print("No changes: the site already has these prices.")
        print_list(kept_note, [names.get(k, k) for k in kept])
        return 0

    # Load the new prices the way the site will, before writing anything.
    with tempfile.TemporaryDirectory() as temp:
        candidate = Path(temp) / "prices.json"
        candidate.write_text(to_json(new) + "\n", encoding="utf-8")
        new_list = read_price_list(candidate)
    expected = {key: (slot["amount"], slot["confirmed"]) for key, slot in price_slots(new).items()}
    expected_removed = sorted(slug for slug, p in new["products"].items() if is_removed(p))
    if site_slots(new_list) != expected or sorted(new_list["removedProducts"]) != expected_removed:
        raise Stop("The site would not show the prices from the spreadsheet (prices.json and src/config/prices.ts disagree). Nothing was changed.")

    if changes:
        print(f"\n{len(changes)} price change{'s' if len(changes) != 1 else ''}:")
        for key, old, slot, why in changes:
            print(describe(names.get(key, key), old, slot, why))
    gone_products = [k for k in removals if k.count("/") == 1]
    gone_sizes = [k for k in removals if k.count("/") == 2 and not k.endswith("/lights")]
    gone_lights = [k for k in removals if k.endswith("/lights")]
    n = len(gone_products)
    print_list(f"{n} piece{'s' if n != 1 else ''} taken off the site (all {'their' if n != 1 else 'its'} rows were deleted):" if n else "", [names.get(k, k) for k in gone_products])
    print_list("Sizes no longer offered (row deleted):", [names.get(k, k) for k in gone_sizes])
    print_list("LED lights add-on switched off (row deleted):", [names.get(k, k) for k in gone_lights])
    print_list(kept_note, [names.get(k, k) for k in kept])

    notes, unconfirmed = [], []
    for key, old, slot, why in changes:
        label = names.get(key, key)
        if old["amount"] and slot["amount"] and (slot["amount"] >= 3 * old["amount"] or 3 * slot["amount"] <= old["amount"]):
            notes.append(f"Big change, please double-check: {label} {inr(old['amount'])} → {inr(slot['amount'])}.")
        if slot["amount"] is None and old["amount"] is not None and why != "deleted":
            notes.append(f"{label} now shows no price (price on request).")
        if slot["amount"] != old["amount"] and not slot["confirmed"]:
            unconfirmed.append(label)
    if unconfirmed:
        notes.append(f"New price{'s' if len(unconfirmed) > 1 else ''} still marked Confirmed = No (fine if still a guess): {', '.join(unconfirmed)}.")
    nav = (ROOT / "src" / "config" / "nav.ts").read_text(encoding="utf-8")
    linked = [names.get(k, k) for k in gone_products if f'"/product/{k.split("/")[1]}"' in nav]
    if linked:
        notes.append(f"Still linked from the header menu (src/config/nav.ts), where the link now leads to “not found”; ask the developer to take it out: {', '.join(linked)}.")
    if notes:
        print("\nNotes:")
        for note in notes:
            print(f"  • {note}")

    if dry_run:
        print("\nDry run: nothing was written.")
        return 0

    # Refresh the main catalogue afterwards only if it holds no other edits of its own.
    refresh = not no_export and (source.resolve() == CATALOGUE.resolve() or not CATALOGUE.exists())
    if not no_export and not refresh:
        other_errors, other_changes, other_removals = compare_workbook(CATALOGUE, data)
        refresh = not other_errors and not other_changes and not other_removals

    temp = PRICES_JSON.with_name(".prices.json.saving")
    temp.write_text(to_json(new) + "\n", encoding="utf-8")
    os.replace(temp, PRICES_JSON)
    print(f"\nSaved {display(PRICES_JSON)}.")
    if no_export:
        print(f"{display(CATALOGUE)} was left as it is (--no-export). Run npm run catalogue:export later so it shows the saved prices.")
    elif refresh:
        try:
            write_catalogue(CATALOGUE, new, new_list, seed)
            print(f"Rebuilt {display(CATALOGUE)} with the saved prices.")
        except Stop as error:
            print(f"{error}\nThen run npm run catalogue:export so the catalogue shows the saved prices.")
    else:
        print(f"{display(CATALOGUE)} has other edits, so it was left as it is.")
    print("Next: check the site (restart npm run preview if it was running, so the cart uses the new prices), then commit and push.")
    return 0


def main():
    parser = argparse.ArgumentParser(description="Art Studio price catalogue (see the top of this file).")
    commands = parser.add_subparsers(dest="command", required=True)
    export = commands.add_parser("export", help=f"write {display(CATALOGUE)}")
    export.add_argument("--force", action="store_true", help="replace the catalogue even if it has edits not imported yet")
    imp = commands.add_parser("import", help="read the catalogue and write src/data/prices.json")
    imp.add_argument("path", nargs="?", default=str(CATALOGUE), help=f"spreadsheet to read (default {display(CATALOGUE)})")
    imp.add_argument("--dry-run", action="store_true", help="check and list the changes without writing anything")
    imp.add_argument("--no-export", action="store_true", help="don't rebuild the catalogue afterwards (e.g. while it is open in Excel)")
    args = parser.parse_args()
    try:
        if args.command == "export":
            command_export(args.force)
            return 0
        return command_import(args.path, args.dry_run, args.no_export)
    except Stop as error:
        print(error, file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
